import io
import logging
import re

from fastapi import (
    APIRouter, File, Form, HTTPException, Query, Request, UploadFile,
)
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from app.config import DEFAULT_PASSWORD, ROLE_NAMES
from app.database import SessionLocal
from app.lookups import _get_user_state_name
from app.security import hash_password, verify_password
from app.templating import templates
from app.utils import next_numeric_id, next_species_id

logger = logging.getLogger("lbvd")
router = APIRouter(tags=["system_admin"])

# Username rules: English letters/digits plus @ and _ only, never starting with a digit.
USERNAME_MAX_LEN = 50
USERNAME_PATTERN = re.compile(r"^[A-Za-z_@][A-Za-z0-9_@]*$")
USERNAME_RULE_MESSAGE = (
    "❌ အသုံးပြုသူအမည်တွင် အင်္ဂလိပ်စာလုံး (A-Z, a-z)၊ အင်္ဂလိပ်ဂဏန်း (0-9)၊ "
    "@ နှင့် _ သာ သုံးနိုင်ပါသည်။ ဂဏန်းဖြင့် စတင်၍လည်း မရပါ။"
)


def _is_valid_username(username: str) -> bool:
    """English letters/digits plus @/_ only, never starting with a digit."""
    if not username or len(username) > USERNAME_MAX_LEN:
        return False
    return bool(USERNAME_PATTERN.fullmatch(username))


def _normalize_username(raw: str) -> str:
    """Trim and validate a username, raising HTTPException on invalid input."""
    username = (raw or "").strip()
    if not username:
        raise HTTPException(400, "❌ အသုံးပြုသူအမည် ထည့်ပါ။")
    if len(username) > USERNAME_MAX_LEN:
        raise HTTPException(
            400,
            f"❌ အသုံးပြုသူအမည်သည် စာလုံး {USERNAME_MAX_LEN} ခုထက် မကျော်ရပါ။",
        )
    if not _is_valid_username(username):
        raise HTTPException(400, USERNAME_RULE_MESSAGE)
    return username


def _username_taken(db, username: str, exclude_uid=None) -> bool:
    """True when the username is already used as a login name or display name."""
    sql = ("SELECT userId FROM `user` "
           "WHERE (LOWER(loginName) = LOWER(:uname) OR LOWER(userName) = LOWER(:uname))")
    params = {"uname": username}
    if exclude_uid:
        sql += " AND userId != :uid"
        params["uid"] = exclude_uid
    return db.execute(text(sql + " LIMIT 1"), params).fetchone() is not None


# ============================================================
# 🏠 SYSTEM DASHBOARD
# ============================================================
@router.get("/system/dashboard", response_class=HTMLResponse)
async def system_dashboard_page(request: Request):
    user = request.session.get("user")
    if not user:
        return RedirectResponse("/", status_code=303)
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ဝင်ရောက်နိုင်ပါသည်။")

    db = SessionLocal()
    role_counts = {"state": 0, "district": 0, "township": 0, "enumerator": 0, "union": 0}
    active_counts = dict(role_counts)
    total_users = total_animals = total_records = 0
    try:
        rows = db.execute(
            text("SELECT userRole, stateID, COUNT(*) AS cnt FROM `user` "
                 "WHERE userRole != 1 GROUP BY userRole, stateID"),
        ).fetchall()
        for row in rows:
            rid = int(row[0]); sid = (row[1] or "").strip(); cnt = int(row[2])
            name = ROLE_NAMES.get(rid)
            if name == "state" and sid == "s15":
                role_counts["union"] += cnt
            elif name in role_counts:
                role_counts[name] += cnt

        active_rows = db.execute(
            text("SELECT userRole, stateID, COUNT(*) AS cnt FROM `user` "
                 "WHERE userRole != 1 AND status = 1 GROUP BY userRole, stateID"),
        ).fetchall()
        for row in active_rows:
            rid = int(row[0]); sid = (row[1] or "").strip(); cnt = int(row[2])
            name = ROLE_NAMES.get(rid)
            if name == "state" and sid == "s15":
                active_counts["union"] += cnt
            elif name in active_counts:
                active_counts[name] += cnt

        total_users = sum(role_counts.values())
        r = db.execute(text("SELECT COUNT(*) FROM animaltype")).fetchone()
        total_animals = int(r[0]) if r else 0
        r = db.execute(text("SELECT COUNT(*) FROM personcv")).fetchone()
        total_records = int(r[0]) if r else 0
    except SQLAlchemyError as e:
        logger.error("system_dashboard error: %s", e)
    finally:
        db.close()

    return templates.TemplateResponse(request=request, name="systemAdminDashboard.html", context={
        "active_page": "sys_home", "user": user,
        "user_state_id": user.get("stateID") or "",
        "user_state_name": _get_user_state_name(user),
        "role_counts": role_counts, "active_counts": active_counts,
        "total_users": total_users, "total_animals": total_animals,
        "total_records": total_records,
    })


# ============================================================
# 👥 USERS
# ============================================================
@router.get("/system/users", response_class=HTMLResponse)
async def system_users_page(request: Request):
    user = request.session.get("user")
    if not user:
        return RedirectResponse("/", status_code=303)
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ဝင်ရောက်နိုင်ပါသည်။")
    return templates.TemplateResponse(request=request, name="systemUserCreate.html", context={
        "active_page": "sys_users",
        "user_state_id": user.get("stateID") or "",
        "user_state_name": _get_user_state_name(user),
    })


@router.get("/api/system/roles")
async def get_system_roles():
    return {"roles": [
        {"id": 1, "name": "admin"}, {"id": 2, "name": "union"},
        {"id": 3, "name": "state"}, {"id": 4, "name": "district"},
        {"id": 5, "name": "township"}, {"id": 6, "name": "enumerator"},
    ]}


@router.get("/api/system/users")
async def api_get_system_users(
    request: Request,
    role_filter: str = Query(""), state: str = Query(""),
    district: str = Query(""), township: str = Query(""),
):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ကြည့်ရှုနိုင်ပါသည်။")
    db = SessionLocal()
    try:
        conds = ["1=1"]; params = {}
        if role_filter: conds.append("u.userRole = :role"); params["role"] = role_filter
        if state:       conds.append("u.stateID = :sid");   params["sid"] = state
        if district:    conds.append("u.districtID = :did"); params["did"] = district
        if township:    conds.append("u.townshipID = :tid"); params["tid"] = township
        where = " AND ".join(conds)
        sql = text(f"""
            SELECT u.userId, u.userName, u.loginName AS email, u.phone,
                   u.userRole, u.status, u.stateID, st.stateName,
                   u.districtID, dt.districtName, u.townshipID, ts.townshipName, ur.roleName
            FROM `user` u
            LEFT JOIN state    st ON st.stateID = u.stateID
            LEFT JOIN district dt ON dt.stateID = u.stateID AND dt.districtID = u.districtID
            LEFT JOIN township ts ON ts.stateID = u.stateID AND ts.districtID = u.districtID
                                 AND ts.townshipID = u.townshipID
            LEFT JOIN userrole ur ON ur.roleId = u.userRole
            WHERE {where} ORDER BY u.userId DESC
        """)
        rows = db.execute(sql, params).mappings().all()
        return {"status": "success", "users": [dict(r) for r in rows]}
    except SQLAlchemyError as e:
        logger.error("System users API error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.get("/api/system/users/search")
async def api_search_system_users(
    request: Request, q: str = Query(""), role_filter: str = Query(""),
    state: str = Query(""), district: str = Query(""), township: str = Query(""),
    page: int = Query(1, ge=1), per_page: int = Query(10, ge=1, le=100),
):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ကြည့်ရှုနိုင်ပါသည်။")
    db = SessionLocal()
    try:
        conds = ["1=1"]; params = {}
        if q and q.strip():
            conds.append("(u.userName LIKE :q OR u.loginName LIKE :q)")
            params["q"] = f"%{q.strip()}%"
        if role_filter: conds.append("u.userRole = :role"); params["role"] = role_filter
        if state:       conds.append("u.stateID = :sid");   params["sid"] = state
        if district:    conds.append("u.districtID = :did"); params["did"] = district
        if township:    conds.append("u.townshipID = :tid"); params["tid"] = township
        where = " AND ".join(conds)

        total = db.execute(text(f"SELECT COUNT(*) AS total FROM `user` u WHERE {where}"),
                           params).scalar() or 0
        offset = (page - 1) * per_page
        sql = text(f"""
            SELECT u.userId, u.userName, u.loginName AS email, u.phone,
                   u.userRole, u.status, u.stateID, st.stateName,
                   u.districtID, dt.districtName, u.townshipID, ts.townshipName, ur.roleName
            FROM `user` u
            LEFT JOIN state    st ON st.stateID = u.stateID
            LEFT JOIN district dt ON dt.stateID = u.stateID AND dt.districtID = u.districtID
            LEFT JOIN township ts ON ts.stateID = u.stateID AND ts.districtID = u.districtID
                                 AND ts.townshipID = u.townshipID
            LEFT JOIN userrole ur ON ur.roleId = u.userRole
            WHERE {where} ORDER BY u.userId DESC LIMIT :limit OFFSET :offset
        """)
        params["limit"] = per_page
        params["offset"] = offset
        rows = db.execute(sql, params).mappings().all()
        total_pages = (total + per_page - 1) // per_page if total > 0 else 1
        return {"status": "success", "users": [dict(r) for r in rows],
                "total": total, "page": page, "per_page": per_page,
                "total_pages": total_pages}
    except SQLAlchemyError as e:
        logger.error("Search users error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/users/action")
async def api_system_user_action(request: Request):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ လုပ်ဆောင်နိုင်ပါသည်။")
    data = await request.json()
    target_uid = data.get("userId"); action = data.get("action"); new_status = data.get("status")
    db = SessionLocal()
    try:
        if action == "toggle":
            db.execute(text("UPDATE `user` SET status = :st WHERE userId = :uid"),
                       {"st": new_status, "uid": target_uid})
        db.commit()
        return {"status": "success", "message": "အောင်မြင်ပါသည်။"}
    except SQLAlchemyError:
        db.rollback()
        raise HTTPException(500, "Database error")
    finally:
        db.close()


@router.get("/api/system/users/check-username")
async def api_check_username(
    request: Request, username: str = Query(""), exclude: int = Query(0),
):
    """Live uniqueness/format check used by the create/edit user form."""
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") not in ("admin", "union"):
        raise HTTPException(403, "Admin သာ စစ်ဆေးနိုင်ပါသည်။")
    try:
        uname = _normalize_username(username)
    except HTTPException as exc:
        return {"valid": False, "available": False, "message": exc.detail}

    db = SessionLocal()
    try:
        if _username_taken(db, uname, exclude or None):
            return {"valid": True, "available": False,
                    "message": "❌ ဤအသုံးပြုသူအမည် အသုံးပြုပြီးသား ဖြစ်ပါသည်။"}
        return {"valid": True, "available": True, "message": "✅ ဤအသုံးပြုသူအမည်ကို သုံးနိုင်ပါသည်။"}
    finally:
        db.close()


@router.post("/api/system/users/create")
async def api_create_system_user(
    request: Request, username: str = Form(...), role_id: int = Form(...),
    phone: str = Form(""),
    stateID: str = Form(""), districtID: str = Form(""), townshipID: str = Form(""),
):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") not in ("admin", "union"):
        raise HTTPException(403, "Admin သာ ဖန်တီးနိုင်ပါသည်။")
    uname = _normalize_username(username)

    db = SessionLocal()
    try:
        if _username_taken(db, uname):
            raise HTTPException(409, "❌ ဤအသုံးပြုသူအမည် အသုံးပြုပြီးသား ဖြစ်ပါသည်။")
        phone_val = phone.strip() if phone else ""
        if phone_val and db.execute(text("SELECT userId FROM `user` WHERE phone = :phone"),
                                    {"phone": phone_val}).fetchone():
            raise HTTPException(409, "❌ ဤဖုန်းနံပါတ် အသုံးပြုပြီးသားပါ။")

        if role_id in (1, 2):
            stateID = districtID = townshipID = None
        elif role_id == 3:
            if not stateID: raise HTTPException(400, "တိုင်းဒေသကြီး ရွေးပါ။")
            districtID = townshipID = None
        elif role_id == 4:
            if not stateID or not districtID:
                raise HTTPException(400, "တိုင်းနှင့် ခရိုင် ရွေးပါ။")
            townshipID = None
        elif role_id in (5, 6):
            if not stateID or not districtID or not townshipID:
                raise HTTPException(400, "တိုင်း၊ ခရိုင်၊ မြို့နယ် အားလုံး ရွေးပါ။")

        db.execute(text("""
            INSERT INTO `user` (userName, userRole, loginName, password,
                stateID, districtID, townshipID, status, phone)
            VALUES (:uname, :role, :login, :pwd, :sid, :did, :tid, 1, :phone)
        """), {
            "uname": uname, "role": role_id, "login": uname,
            "pwd": hash_password(DEFAULT_PASSWORD),
            "sid": stateID or None, "did": districtID or None,
            "tid": townshipID or None,
            "phone": phone_val if phone_val else None,
        })
        db.commit()
        return {"status": "success",
                "message": "အကောင့်အသစ် ဖန်တီးပြီးပါပြီ။ (လျှို့ဝှက်နံပါတ်: 12345678)"}
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "❌ အချက်အလက်တူညီနေပါသည်။")
    except SQLAlchemyError as e:
        db.rollback()
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/users/update")
async def api_update_system_user(
    request: Request, userId: int = Form(...), username: str = Form(...),
    role_id: int = Form(...), phone: str = Form(""),
    stateID: str = Form(""), districtID: str = Form(""), townshipID: str = Form(""),
):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") not in ("admin", "union"):
        raise HTTPException(403, "Admin သာ ပြင်ဆင်နိုင်ပါသည်။")
    raw_username = (username or "").strip()
    if not raw_username:
        raise HTTPException(400, "❌ အသုံးပြုသူအမည် ထည့်ပါ။")

    db = SessionLocal()
    try:
        current = db.execute(
            text("SELECT userName, loginName FROM `user` WHERE userId = :uid"),
            {"uid": userId},
        ).mappings().first()
        if not current:
            raise HTTPException(404, "အသုံးပြုသူ မတွေ့ပါ။")

        # Accounts created before the username-login change keep their existing
        # display name / login name when the username field is left untouched.
        username_changed = raw_username != str(current["userName"] or "").strip()
        if username_changed:
            uname = _normalize_username(raw_username)
            if _username_taken(db, uname, exclude_uid=userId):
                raise HTTPException(409, "❌ ဤအသုံးပြုသူအမည်ကို အခြားသူတစ်ဦး အသုံးပြုနေပါသည်။")
            login_name = uname
        else:
            uname = raw_username
            login_name = current["loginName"]

        phone_val = phone.strip() if phone else ""
        if phone_val and db.execute(text("SELECT userId FROM `user` "
                                         "WHERE phone = :phone AND userId != :uid"),
                                    {"phone": phone_val, "uid": userId}).fetchone():
            raise HTTPException(409, "❌ ဤဖုန်းနံပါတ်ကို အခြားသူတစ်ဦး အသုံးပြုနေပါသည်။")

        if role_id in (1, 2):
            stateID = districtID = townshipID = None
        elif role_id == 3:
            if not stateID: raise HTTPException(400, "တိုင်းဒေသကြီး ရွေးပါ။")
            districtID = townshipID = None
        elif role_id == 4:
            if not stateID or not districtID:
                raise HTTPException(400, "တိုင်းနှင့် ခရိုင် ရွေးပါ။")
            townshipID = None
        elif role_id in (5, 6):
            if not stateID or not districtID or not townshipID:
                raise HTTPException(400, "တိုင်း၊ ခရိုင်၊ မြို့နယ် အားလုံး ရွေးပါ။")

        db.execute(text("""
            UPDATE `user` SET userName = :uname, userRole = :role, loginName = :login,
                stateID = :sid, districtID = :did, townshipID = :tid, phone = :phone
            WHERE userId = :uid
        """), {
            "uname": uname, "role": role_id, "login": login_name,
            "sid": stateID or None, "did": districtID or None,
            "tid": townshipID or None,
            "phone": phone_val if phone_val else None, "uid": userId,
        })
        db.commit()
        return {"status": "success", "message": "အချက်အလက်များ ပြင်ဆင်ပြီးပါပြီ။"}
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "❌ အချက်အလက်တူညီနေပါသည်။")
    except SQLAlchemyError as e:
        db.rollback()
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/users/import-excel")
async def api_import_users_excel(request: Request, file: UploadFile = File(...)):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") not in ("admin", "union"):
        raise HTTPException(403, "Admin သာ လုပ်ဆောင်နိုင်ပါသည်။")
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "❌ ကျေးဇူးပြု၍ .xlsx သို့မဟုတ် .xls ဖိုင်သာ တင်ပါ။")

    hashed_pwd = hash_password(DEFAULT_PASSWORD)
    db = SessionLocal()
    success_count = 0
    failed_rows = []
    try:
        contents = await file.read()
        wb = load_workbook(filename=io.BytesIO(contents), data_only=True)
        if len(wb.sheetnames) > 1:
            raise HTTPException(400, "❌ Sheet တစ်ခုထက်ပို၍ ပါဝင်နေပါသည်။")
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(h).strip() for h in header_row if h is not None]
        required_cols = ["userName", "userRole"]
        missing_cols = [col for col in required_cols if col not in headers]
        if missing_cols:
            raise HTTPException(400, f"❌ လိုအပ်သော Column များ ပျောက်နေပါသည်: {', '.join(missing_cols)}")

        idx_name = headers.index("userName")
        idx_role = headers.index("userRole")
        idx_state = headers.index("stateID") if "stateID" in headers else -1
        idx_district = headers.index("districtID") if "districtID" in headers else -1
        idx_township = headers.index("townshipID") if "townshipID" in headers else -1
        idx_phone = headers.index("phone") if "phone" in headers else -1

        def get_val(row, idx):
            if idx == -1 or idx >= len(row):
                return ""
            val = row[idx]
            return str(val).strip() if val is not None else ""

        seen_usernames = set()

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(cell is None for cell in row):
                continue
            username = get_val(row, idx_name)
            role_str = get_val(row, idx_role)
            state_id = get_val(row, idx_state)
            district_id = get_val(row, idx_district)
            township_id = get_val(row, idx_township)
            phone = get_val(row, idx_phone)

            try:
                role_id = int(float(role_str)) if role_str else 0
            except ValueError:
                role_id = 0
            if role_id not in (1, 2, 3, 4, 5, 6):
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "Role မမှန်ကန်ပါ။ (1-6 အတွင်းသာ)"})
                continue
            if role_id == 3 and not state_id:
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "Role 3 အတွက် stateID လိုအပ်ပါသည်။"})
                continue
            if role_id == 4 and (not state_id or not district_id):
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "Role 4 အတွက် stateID+districtID လိုအပ်ပါသည်။"})
                continue
            if role_id in (5, 6) and (not state_id or not district_id or not township_id):
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "Role 5,6 အတွက် state+district+township လိုအပ်ပါသည်။"})
                continue
            if not username:
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "အသုံးပြုသူအမည် မပါဝင်ပါ။"})
                continue
            try:
                username = _normalize_username(username)
            except HTTPException as exc:
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": exc.detail})
                continue
            username_key = username.casefold()
            if username_key in seen_usernames:
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "ဤအသုံးပြုသူအမည်သည် ဖိုင်ထဲတွင် ထပ်နေပါသည်။"})
                continue
            if _username_taken(db, username):
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "ဤအသုံးပြုသူအမည် အသုံးပြုပြီးသား ဖြစ်ပါသည်။"})
                continue
            if phone and not re.match(r"^(09|\+959)\d{7,9}$", phone):
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "ဖုန်းနံပါတ် Format မှားယွင်းနေပါသည်။"})
                continue
            if phone and db.execute(text("SELECT userId FROM `user` WHERE phone = :phone"),
                                    {"phone": phone}).fetchone():
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "ဤဖုန်းနံပါတ် အသုံးပြုပြီးသား ဖြစ်ပါသည်။"})
                continue
            seen_usernames.add(username_key)

            final_sid = final_did = final_tid = None
            if role_id == 3:
                final_sid = state_id
            elif role_id == 4:
                final_sid, final_did = state_id, district_id
            elif role_id in (5, 6):
                final_sid, final_did, final_tid = state_id, district_id, township_id

            try:
                db.execute(text("""
                    INSERT INTO `user` (userName, userRole, loginName, password,
                        stateID, districtID, townshipID, status, phone)
                    VALUES (:uname, :role, :login, :pwd, :sid, :did, :tid, 1, :phone)
                """), {
                    "uname": username, "role": role_id, "login": username,
                    "pwd": hashed_pwd,
                    "sid": final_sid or None, "did": final_did or None,
                    "tid": final_tid or None,
                    "phone": phone if phone else None,
                })
                db.commit()
                success_count += 1
            except IntegrityError:
                db.rollback()
                failed_rows.append({"row": row_idx, "username": username,
                                    "reason": "Database တွင် အချက်အလက် ထပ်နေပါသည်။"})

        return {"status": "success",
                "message": "အောင်မြင်စွာ ထည့်သွင်းပြီးပါပြီ။",
                "success_count": success_count,
                "duplicate_count": len(failed_rows),
                "failed_rows": failed_rows,
                "default_password": DEFAULT_PASSWORD}
    except HTTPException:
        db.rollback(); raise
    except Exception as e:
        db.rollback()
        raise HTTPException(500, f"❌ Excel ဖတ်ရာတွင် အမှားအယွင်း ရှိနေပါသည်။ ({str(e)})")
    finally:
        db.close()


@router.get("/api/system/users/export-data")
async def api_export_users_data(request: Request, roles: str = Query("")):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ လုပ်ဆောင်နိုင်ပါသည်။")
    if not roles:
        return {"status": "success", "users": []}
    role_list = [r.strip() for r in roles.split(",") if r.strip()]
    if not role_list:
        return {"status": "success", "users": []}

    db = SessionLocal()
    try:
        placeholders = ", ".join([f":r{i}" for i in range(len(role_list))])
        params = {f"r{i}": r for i, r in enumerate(role_list)}
        sql = text(f"""
            SELECT u.userId, u.userName, u.userRole,
                   u.stateID, st.stateName, u.districtID, dt.districtName,
                   u.townshipID, ts.townshipName, u.phone, u.status
            FROM `user` u
            LEFT JOIN state    st ON st.stateID = u.stateID
            LEFT JOIN district dt ON dt.stateID = u.stateID AND dt.districtID = u.districtID
            LEFT JOIN township ts ON ts.stateID = u.stateID AND ts.districtID = u.districtID
                                 AND ts.townshipID = u.townshipID
            WHERE u.userRole IN ({placeholders}) ORDER BY u.userRole, u.userName
        """)
        rows = db.execute(sql, params).mappings().all()
        return {"status": "success", "users": [dict(r) for r in rows]}
    except SQLAlchemyError as e:
        logger.error("Export users data error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/users/import-excel/errors")
async def api_download_error_excel(request: Request):
    data = await request.json()
    errors = data.get("errors", [])
    if not errors:
        raise HTTPException(400, "❌ Download လုပ်ရန် Error မရှိပါ။")

    wb = Workbook()
    ws = wb.active
    ws.title = "Import Errors"
    ws.append(["Row နံပါတ်", "အသုံးပြုသူအမည်", "အကြောင်းရင်း"])
    for err in errors:
        ws.append([err.get("row"), err.get("username"),
                   err.get("reason") or "အမှားအယွင်း ရှိနေပါသည်။"])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 60

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Import_Errors.xlsx"},
    )


@router.post("/api/system/users/reset-password")
async def api_reset_user_password(request: Request):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ လုပ်ဆောင်နိုင်ပါသည်။")
    try:
        data = await request.json()
    except Exception:
        raise HTTPException(400, "Invalid JSON body.")
    target_uid = data.get("userId")
    if not target_uid:
        raise HTTPException(400, "userId လိုအပ်ပါသည်။")

    db = SessionLocal()
    try:
        existing = db.execute(text("SELECT userId, userName FROM `user` WHERE userId = :uid"),
                              {"uid": target_uid}).fetchone()
        if not existing:
            raise HTTPException(404, "အသုံးပြုသူ မတွေ့ပါ။")
        db.execute(text("UPDATE `user` SET password = :pwd WHERE userId = :uid"),
                   {"pwd": hash_password(DEFAULT_PASSWORD), "uid": target_uid})
        db.commit()
        return {"status": "success",
                "message": f"'{existing[1]}' ရဲ့ Password ကို default ({DEFAULT_PASSWORD}) ပြန်ပြောင်းပြီးပါပြီ။",
                "default_password": DEFAULT_PASSWORD}
    except HTTPException:
        db.rollback(); raise
    except SQLAlchemyError as e:
        db.rollback()
        logger.error("Reset password error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


# ============================================================
# 🐄 ANIMALS
# ============================================================
@router.get("/system/animals", response_class=HTMLResponse)
async def system_animals_page(request: Request):
    user = request.session.get("user")
    if not user:
        return RedirectResponse("/", status_code=303)
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ဝင်ရောက်နိုင်ပါသည်။")
    return templates.TemplateResponse(request=request, name="systemAnimals.html", context={
        "active_page": "sys_animals",
        "user_state_id": user.get("stateID") or "",
        "user_state_name": _get_user_state_name(user),
    })


@router.post("/api/system/animals/create")
async def api_create_animal(request: Request):
    admin_user = request.session.get("user")
    if not admin_user or admin_user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ဖန်တီးနိုင်ပါသည်။")
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(400, "Invalid JSON body.")

    category = (payload.get("category") or "").strip()
    animal_name = (payload.get("animalName") or "").strip()
    needs_breed = bool(payload.get("needsBreed"))
    raw_breeds = payload.get("breeds") or []

    if not animal_name:
        raise HTTPException(400, "တိရစ္ဆာန်အမည် ထည့်ပါ။")
    if len(animal_name) > 30:
        raise HTTPException(400, "တိရစ္ဆာန်အမည်သည် ၃၀ လုံးထက် မကျော်ရပါ။")
    if not isinstance(raw_breeds, list):
        raise HTTPException(400, "မျိုးစိတ်စာရင်း ပုံစံ မမှန်ကန်ပါ။")

    cleaned_breeds = []
    seen_breeds = set()
    for b in raw_breeds:
        b = str(b).strip()
        if not b:
            continue
        if len(b) > 30:
            raise HTTPException(400, f"မျိုးစိတ်အမည် '{b[:20]}...' သည် ၃၀ လုံးထက် ကျော်နေပါသည်။")
        key = b.lower()
        if key in seen_breeds:
            continue
        seen_breeds.add(key)
        cleaned_breeds.append(b)
    if needs_breed and not cleaned_breeds:
        raise HTTPException(400, "မျိုးစိတ် အနည်းဆုံး တစ်ခု ထည့်ပါ။")

    CATEGORY_MAP = {
        "large":   ("atID01", "largeanimals",  "largespecies",     "speciesName", "la", "laS"),
        "small":   ("atID02", "younganimals",  "youngspecies",     "speciesName", "ya", "yaS"),
        "poultry": ("atID03", "livestock",     "livestockspecies", "speciesName", "ls", "lsS"),
        "breeder": ("atID04", "breedingstock", None,                None,         "bs", None),
    }
    if category not in CATEGORY_MAP:
        raise HTTPException(400, "အမျိုးအစား မမှန်ကန်ပါ။")

    animal_type_id, animal_table, species_table, species_name_col, animal_prefix, species_prefix = \
        CATEGORY_MAP[category]

    db = SessionLocal()
    try:
        existing = db.execute(
            text(f"SELECT animalID FROM `{animal_table}` "
                 "WHERE animalTypeID = :tid AND animalName = :name"),
            {"tid": animal_type_id, "name": animal_name},
        ).fetchone()
        if existing:
            raise HTTPException(409, "ဤတိရစ္ဆာန်အမည် ရှိပြီးသားပါ။")

        new_animal_id = next_numeric_id(db, animal_table, "animalID", animal_prefix, 2)
        db.execute(text(f"INSERT INTO `{animal_table}` "
                        "(animalTypeID, animalID, animalName) VALUES (:tid, :aid, :name)"),
                   {"tid": animal_type_id, "aid": new_animal_id, "name": animal_name})

        inserted_breeds = []
        if needs_breed and species_table:
            for breed in cleaned_breeds:
                new_species_id = next_species_id(db, species_table, species_prefix,
                                                 new_animal_id, animal_prefix)
                db.execute(
                    text(f"INSERT INTO `{species_table}` "
                         f"(animalID, speciesID, `{species_name_col}`) "
                         "VALUES (:aid, :sid, :name)"),
                    {"aid": new_animal_id, "sid": new_species_id, "name": breed},
                )
                inserted_breeds.append({"speciesID": new_species_id, "name": breed})

        db.commit()
        return {"status": "success",
                "message": "တိရစ်ဆာန်အမျိုးအစား ထည့်သွင်းပြီးပါပြီ။",
                "animalID": new_animal_id, "animalName": animal_name,
                "animalTypeID": animal_type_id, "breeds": inserted_breeds,
                "breedsInserted": len(inserted_breeds)}
    except HTTPException:
        db.rollback(); raise
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "အချက်အလက် ထပ်နေပါသည်။")
    except SQLAlchemyError as e:
        db.rollback()
        raise HTTPException(500, f"Database error: {e}")
    except Exception as e:
        db.rollback()
        raise HTTPException(500, f"Unexpected error: {e}")
    finally:
        db.close()


# ============================================================
# 👤 ADMIN PROFILE
# ============================================================
@router.get("/system/profile", response_class=HTMLResponse)
async def system_profile_page(request: Request):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ဝင်ရောက်နိုင်ပါသည်။")
    return templates.TemplateResponse(request=request, name="systemAdminProfile.html", context={
        "active_page": "sys_profile",
        "user_state_id": user.get("stateID") or "",
        "user_state_name": _get_user_state_name(user),
    })


@router.get("/api/system/profile")
async def api_get_admin_profile(request: Request):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ကြည့်ရှုနိုင်ပါသည်။")
    db = SessionLocal()
    try:
        row = db.execute(
            text("SELECT userId, userName, loginName, phone, userRole, status "
                 "FROM `user` WHERE userId = :uid"),
            {"uid": user["id"]},
        ).mappings().first()
        if not row:
            raise HTTPException(404, "အကောင့် မတွေ့ပါ။")
        return {"status": "success", "profile": dict(row)}
    except HTTPException:
        raise
    except SQLAlchemyError as e:
        logger.error("Get profile error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/profile/update-info")
async def api_update_admin_info(
    request: Request, username: str = Form(...), email: str = Form(...),
    phone: str = Form(""), old_credential: str = Form(""),
):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ ပြင်ဆင်နိုင်ပါသည်။")
    db = SessionLocal()
    try:
        current = db.execute(
            text("SELECT userId, userName, loginName, password "
                 "FROM `user` WHERE userId = :uid"),
            {"uid": user["id"]},
        ).mappings().first()
        if not current:
            raise HTTPException(404, "အကောင့် မတွေ့ပါ။")

        changed = (str(current["userName"]).strip() != username.strip()
                   or str(current["loginName"]).strip() != email.strip())
        if changed:
            if not old_credential.strip():
                raise HTTPException(400, "အသုံးပြုသူအမည် သို့မဟုတ် အီးမေးလ် ပြောင်းရန် "
                                         "လက်ရှိ Login သို့မဟုတ် Password ကို ရိုက်ထည့်ပါ။")
            cred = old_credential.strip()
            if not (cred == str(current["loginName"]).strip()
                    or verify_password(cred, current["password"])):
                raise HTTPException(401, "လက်ရှိ အီးမေးလ် သို့မဟုတ် စကားဝှက် မှားနေပါသည်။")

        if str(current["userName"]).strip() != username.strip():
            if db.execute(text("SELECT userId FROM `user` "
                               "WHERE userName = :u AND userId != :uid"),
                          {"u": username.strip(), "uid": user["id"]}).fetchone():
                raise HTTPException(409, "❌ ဤအသုံးပြုသူအမည်ကို အခြားသူတစ်ဦး အသုံးပြုနေပါသည်။")
        if str(current["loginName"]).strip() != email.strip():
            if db.execute(text("SELECT userId FROM `user` "
                               "WHERE loginName = :e AND userId != :uid"),
                          {"e": email.strip(), "uid": user["id"]}).fetchone():
                raise HTTPException(409, "❌ ဤအီးမေးလ်ကို အခြားသူတစ်ဦး အသုံးပြုနေပါသည်။")
        if phone.strip():
            if db.execute(text("SELECT userId FROM `user` "
                               "WHERE phone = :p AND userId != :uid"),
                          {"p": phone.strip(), "uid": user["id"]}).fetchone():
                raise HTTPException(409, "❌ ဤဖုန်းနံပါတ်ကို အခြားသူတစ်ဦး အသုံးပြုနေပါသည်။")

        db.execute(text("UPDATE `user` SET userName = :uname, loginName = :login, "
                        "phone = :phone WHERE userId = :uid"),
                   {"uname": username.strip(), "login": email.strip(),
                    "phone": phone.strip() or None, "uid": user["id"]})
        db.commit()
        request.session["user"]["loginName"] = email.strip()
        return {"status": "success", "message": "ကိုယ်ရေးအချက်အလက် ပြင်ဆင်ပြီးပါပြီ။"}
    except HTTPException:
        db.rollback(); raise
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "❌ အချက်အလက်တူညီနေပါသည်။")
    except SQLAlchemyError as e:
        db.rollback()
        logger.error("Update admin info error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()


@router.post("/api/system/profile/change-password")
async def api_change_admin_password(
    request: Request, old_password: str = Form(...),
    new_password: str = Form(...), confirm_password: str = Form(...),
):
    user = request.session.get("user")
    if not user or user.get("role") != "admin":
        raise HTTPException(403, "Admin သာ လုပ်ဆောင်နိုင်ပါသည်။")
    if not old_password or not new_password or not confirm_password:
        raise HTTPException(400, "အချက်အလက်အားလုံး ဖြည့်ပါ။")
    if new_password != confirm_password:
        raise HTTPException(400, "Password အသစ်နှစ်ခု မတူညီပါ။")
    if len(new_password) < 8:
        raise HTTPException(400, "Password အသစ်သည် အနည်းဆုံး ၈ လုံး ရှိရပါမည်။")
    if new_password == old_password:
        raise HTTPException(400, "Password အသစ်နှင့် အဟောင်း မတူညီရပါ။")

    db = SessionLocal()
    try:
        row = db.execute(text("SELECT password FROM `user` WHERE userId = :uid"),
                         {"uid": user["id"]}).fetchone()
        if not row:
            raise HTTPException(404, "အကောင့် မတွေ့ပါ။")
        if not verify_password(old_password, row[0]):
            raise HTTPException(401, "လျှို့ဝှက်နံပါတ် အဟောင်း မှားနေပါသည်။")
        db.execute(text("UPDATE `user` SET password = :pwd WHERE userId = :uid"),
                   {"pwd": hash_password(new_password), "uid": user["id"]})
        db.commit()
        return {"status": "success", "message": "လျှို့ဝှက်နံပါတ် ပြောင်းပြီးပါပြီ။"}
    except HTTPException:
        db.rollback(); raise
    except SQLAlchemyError as e:
        db.rollback()
        logger.error("Change password error: %s", e)
        raise HTTPException(500, f"Database error: {e}")
    finally:
        db.close()