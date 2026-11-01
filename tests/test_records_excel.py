import base64
import json
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from itsdangerous import TimestampSigner
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.pool import StaticPool
from starlette.middleware.sessions import SessionMiddleware

from app.database import SQLAlchemyConnection
from app.records_excel import TEMPLATE, build_records_workbook
from app.routers.households import router


def animal(kind, animal_id, breed, age, sex, count):
    return dict(formID="form1", animalTypeID=kind, animalID=animal_id,
                speciesID=breed, species_name=breed, ageID=age, age_name=age,
                genderID=sex, gender_name=sex, data_count=count, animal_name=animal_id)


class WorkbookTests(unittest.TestCase):
    def test_sample_layout_breeds_age_sex_totals_and_literal_text(self):
        household = dict(formID="form1", respondent_name="=1+1", householdNo="00102",
                         respondent_phone="091111111", interviewDate="2026-12-12")
        details = [
            animal("atID01", "laID01", "breed A", "ageID03", "g03", 1),
            animal("atID01", "laID01", "breed A", "ageID01", "g01", 2),
            animal("atID01", "laID01", "breed B", "ageID03", "g02", 1),
            animal("atID01", "laID01", "breed B", "ageID02", "g02", 1),
            animal("atID01", "laID05", "buffalo", "ageID02", "g01", 4),
            animal("atID02", "yaID01", "goat", "ageID05", "g02", 5),
            animal("atID03", "lsID03", "chicken A", "ageID10", "g01", 10),
            animal("atID03", "lsID03", "chicken A", "ageID11", "g03", 3),
            animal("atID03", "lsID03", "chicken B", "ageID11", "g01", 2),
            animal("atID03", "lsID03", "chicken B", "ageID11", "g03", 2),
            animal("atID03", "lsID07", "turkey", "ageID17", "g03", 7),
            animal("atID03", "lsID02", "broiler", "all", "g03", 9),
            animal("atID04", "bsID07", "", "", "g01", 6),
            animal("atID04", "bsID10", "", "", "g03", 8),
        ]
        workbook = load_workbook(build_records_workbook([household, {"formID": "form2"}], details))
        sheet = workbook.active
        self.assertEqual([sheet.cell(row, 1).value for row in range(6, 9)], [1, 2, 3])
        template = load_workbook(TEMPLATE).active
        self.assertEqual(sheet.max_column, 255)
        self.assertEqual(set(map(str, sheet.merged_cells)), set(map(str, template.merged_cells)))
        for row in template.iter_rows(max_row=5):
            for cell in row:
                actual = sheet[cell.coordinate]
                self.assertEqual(actual.value, cell.value)
                self.assertEqual(actual._style, cell._style)
        expected = {"B6": "form1", "H6": "00102", "J6": "=1+1", "M6": "091111111",
                    "T6": "breed A", "W6": 1, "AA6": 2, "AD6": 3,
                    "T7": "breed B", "V7": 1, "Y7": 1, "AD7": 2,
                    "CT6": None, "CW6": 4, "DC6": 4, "DI6": 5, "DN6": 5,
                    "FR6": 10, "FT6": 3, "FU6": 13, "FS7": 2, "FT7": 2, "FU7": 4,
                    "GT6": 7, "GU6": 7, "FP6": 9, "ID6": 6, "ID7": None}
        for address, value in expected.items():
            self.assertEqual(sheet[address].value, value, address)
        self.assertEqual(sheet["J6"].data_type, "s")
        self.assertEqual(sheet["I6"].value.year, 2026)
        raw = workbook["Animal details"]
        self.assertEqual(raw.max_row, len(details) + 1)
        self.assertEqual(sum(row[10] for row in raw.iter_rows(min_row=2, values_only=True)),
                         sum(detail["data_count"] for detail in details))

    def test_household_without_animals_has_one_row(self):
        sheet = load_workbook(build_records_workbook([{"formID": "empty"}], [])).active
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet["B6"].value, "empty")
        self.assertEqual(sheet["S6"].value, "မမွေး")
        self.assertIsNone(sheet["T6"].value)


class ExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        schemas = {
            "personcv": "formID householdNo name gender age eduStatus phoneNo interviewDate headName headGender headAge headEduStatus headPhoneNo Interviewer interviewerPhoneNo WardName stateID districtID townshipID QVID",
            "state": "stateID stateName", "district": "stateID districtID districtName",
            "township": "stateID districtID townshipID townshipName",
            "qv": "stateID districtID townshipID QVID QVName",
            "collecteddata": "formID animalTypeID animalID speciesID ageID genderID dataCount",
            "largeanimals": "animalID animalName", "younganimals": "animalID animalName",
            "livestock": "animalID animalName", "breedingstock": "animalID animalName",
            "largespecies": "speciesID speciesName animalID", "youngspecies": "speciesID speciesName animalID",
            "livestockspecies": "speciesID speciesName animalID", "age": "ageID ageName",
            "agetype": "ageID ageName animalTypeID animalID", "gender": "genderID genderName",
        }
        with self.engine.begin() as connection:
            for table, columns in schemas.items():
                connection.execute(text(f'CREATE TABLE {table} (' + ','.join(f'{column} TEXT' for column in columns.split()) + ')'))
            connection.execute(text("INSERT INTO personcv (formID,stateID,districtID,townshipID,QVID) VALUES ('allowed','s','d','t','q'),('other','s','d2','t','q')"))
            # Composite location IDs must not multiply household rows.
            connection.execute(text("INSERT INTO district VALUES ('s','d','correct'),('elsewhere','d','incorrect')"))
        app = FastAPI()
        app.add_middleware(SessionMiddleware, secret_key="test-key", session_cookie="test-session")
        app.include_router(router)
        self.client = TestClient(app)
        self.database_patch = patch("app.routers.households.get_db_connection", side_effect=lambda: SQLAlchemyConnection(self.engine.connect()))
        self.database_patch.start()

    def tearDown(self):
        self.database_patch.stop()
        self.client.close()
        self.engine.dispose()

    def sign_in(self, role="district"):
        data = base64.b64encode(json.dumps({"user": {"role": role, "stateID": "s", "districtID": "d"}}).encode())
        self.client.cookies.set("test-session", TimestampSigner("test-key").sign(data).decode())

    def test_export_scopes_selected_ids_deduplicates_and_preserves_order(self):
        self.sign_in()
        response = self.client.post("/api/admin/records/export", json=["other", "allowed", "allowed"])
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
        from io import BytesIO
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet["B6"].value, "allowed")
        self.assertEqual(sheet["D6"].value, "correct")
        self.assertIn("attachment;", response.headers["content-disposition"])
        self.sign_in("admin")
        response = self.client.post("/api/admin/records/export", json=["other", "allowed"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual([sheet["B6"].value, sheet["B7"].value], ["other", "allowed"])

    def test_auth_empty_selection_and_inaccessible_selection(self):
        self.assertEqual(self.client.post("/api/admin/records/export", json=["allowed"]).status_code, 401)
        self.sign_in("enumerator")
        self.assertEqual(self.client.post("/api/admin/records/export", json=["allowed"]).status_code, 403)
        self.sign_in()
        self.assertEqual(self.client.post("/api/admin/records/export", json=[]).status_code, 400)
        self.assertEqual(self.client.post("/api/admin/records/export", json=["other"]).status_code, 404)


if __name__ == "__main__":
    unittest.main()
